import io
import re
from pathlib import Path

from docx import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import KnowledgeChunk, KnowledgeDocument


ALLOWED_EXTENSIONS = {".txt", ".md", ".docx"}


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("支持 .txt、.md、.docx 文件；旧版 .doc 请先另存为 .docx")
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        content = "\n".join(paragraphs)
    else:
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("gb18030")
    content = content.strip()
    if not content:
        raise ValueError("文档没有可提取的文字")
    return content


def split_text(content: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120)
    return splitter.split_text(content)


def _terms(text: str) -> set[str]:
    words = set(re.findall(r"[a-zA-Z0-9]{2,}", text.lower()))
    chinese = re.findall(r"[\u4e00-\u9fff]+", text)
    for section in chinese:
        words.update(section[i : i + 2] for i in range(len(section) - 1))
    return words


def retrieve(db: Session, user_id: int, query: str, limit: int = 4) -> list[dict]:
    query_terms = _terms(query)
    if not query_terms:
        return []
    rows = db.execute(
        select(KnowledgeChunk, KnowledgeDocument.filename)
        .join(KnowledgeDocument)
        .where(KnowledgeDocument.user_id == user_id)
    ).all()
    ranked = []
    for chunk, filename in rows:
        terms = _terms(chunk.content)
        score = len(query_terms & terms) / max(len(query_terms), 1)
        if score:
            ranked.append((score, chunk.id, filename, chunk.content))
    ranked.sort(reverse=True)
    return [
        {"document": filename, "chunk_id": chunk_id, "content": content[:1100]}
        for _, chunk_id, filename, content in ranked[:limit]
    ]
